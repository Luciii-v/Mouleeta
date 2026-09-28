from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Title
title = doc.add_heading('MOULEETA V.2: Production Launch Roadmap', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1. Aim
doc.add_heading('1. Aim & Vision', level=1)
doc.add_paragraph(
    'The primary aim is to launch MOULEETA V.2 as a fully functional, premium, production-ready '
    'e-commerce platform. The goal is to provide a flawless, luxury user experience where every '
    'interaction—from browsing to checkout to post-purchase tracking—feels secure, instantaneous, and bespoke. '
    'We are transitioning from a development/sandbox state into a hardened production state.'
)

# 2. Expectations
doc.add_heading('2. Core Expectations', level=1)
p2 = doc.add_paragraph()
p2.add_run('• Security & Payments: ').bold = True
p2.add_run('All transactions must be processed via live keys (Razorpay). No test keys should remain in the production environment.\n')
p2.add_run('• Data Integrity: ').bold = True
p2.add_run('Returns, tracking, and order data must sync perfectly with Shopify and Shiprocket. Mock data and simulated tracking are strictly prohibited in the live environment.\n')
p2.add_run('• User Privacy: ').bold = True
p2.add_run('Authentication (OTPs, NextAuth) must be completely secure. Sandbox mode must remain disabled in production so sensitive information is never exposed.\n')
p2.add_run('• UI/UX Impeccability: ').bold = True
p2.add_run('Only production-ready components (ProductCard, ProductDetail) will be served to the end user. The interface must remain performant, responsive, and visually luxurious.')

# 3. Plan of Action (POA)
doc.add_heading('3. Plan of Action (POA)', level=1)

# Phase 1
p3 = doc.add_paragraph()
p3.add_run('Phase 1: Production Hardening (Immediate)').bold = True
doc.add_paragraph('• Insert actual Live Razorpay Keys into the production environment.\n'
                  '• Verify the newly integrated Shopify GraphQL Admin API for the returns flow.\n'
                  '• Ensure the Shiprocket API token is valid and live tracking data flows correctly.\n'
                  '• Confirm the Shopify install callback URI functions correctly on the live domain (mouleeta.shop).', style='List Bullet')

# Phase 2
p4 = doc.add_paragraph()
p4.add_run('Phase 2: End-to-End (E2E) Live Testing (Next 3-5 Days)').bold = True
doc.add_paragraph('• Conduct "Penny Testing" to verify Razorpay checkout captures and records payments in Shopify.\n'
                  '• Submit a live return request to verify it appears in the Shopify Admin panel.\n'
                  '• Place a test order and monitor Shiprocket tracking status updates through the frontend UI.\n'
                  '• Perform QA on all core pages (Shop, Collections, Products) to ensure no sandbox UI remains.', style='List Bullet')

# Phase 3
p5 = doc.add_paragraph()
p5.add_run('Phase 3: Soft Launch & Analytics (Next 1-2 Weeks)').bold = True
doc.add_paragraph('• Roll out the platform to a restricted audience (friends, family, VIPs).\n'
                  '• Monitor server logs (Vercel) for unexpected 500 errors or failed API integrations.\n'
                  '• Audit Core Web Vitals (LCP, INP) to ensure the site loads quickly on mobile networks.', style='List Bullet')

# Phase 4
p6 = doc.add_paragraph()
p6.add_run('Phase 4: Full Public Launch').bold = True
doc.add_paragraph('• Execute the full public launch marketing campaign.\n'
                  '• Maintain 24/7 monitoring for the first 48 hours to resolve any edge-case bugs.', style='List Bullet')

# Save Document
doc.save('MOULEETA_V2_Roadmap.docx')
print("Document saved successfully.")
